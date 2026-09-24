"""Fill Template_Tec.pptx with OMX Framework research content."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation

SRC = Path(r"C:\Users\JAVIER YAHIR\Downloads\Template_Tec.pptx")
OUT = Path(r"C:\Users\JAVIER YAHIR\Documents\GitHub\omx-framework\docs\OMX_Framework_Research_Progress.pptx")
OUT_ALT = Path(r"C:\Users\JAVIER YAHIR\Downloads\Template_Tec_OMX_Framework.pptx")

TODAY = date.today().strftime("%Y/%m/%d")


def set_shape_text(shape, text: str) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text


def shapes_with_text(slide):
    return [s for s in slide.shapes if getattr(s, "has_text_frame", False) and s.text_frame]


def apply_slide_texts(slide, texts: list[str]) -> None:
    shapes = shapes_with_text(slide)
    for i, shape in enumerate(shapes):
        txt = texts[i] if i < len(texts) else ""
        set_shape_text(shape, txt)


def main() -> None:
    prs = Presentation(str(SRC))

    # Slide 1 (2 text shapes)
    apply_slide_texts(
        prs.slides[0],
        [
            f"JAVIER YAHIR\nOMX Framework — Agente ROS desacoplado\nDate: {TODAY}",
            "OMX Framework: clasificación de residuos y análisis de estanterías en kombini (vida real)",
        ],
    )

    # Slide 2
    apply_slide_texts(prs.slides[1], ["Problema de investigación y motivación"])

    # Slide 3 (3 shapes: title area, subtitle, body)
    apply_slide_texts(
        prs.slides[2],
        [
            "",
            "Problema social: residuos y gestión de estanterías en tiendas tipo kombini",
            """En Japón y contextos urbanos densos, los kombini deben separar residuos correctamente y mantener estanterías ordenadas con poco personal.

¿Qué problema?
• Clasificación errática de envases/residuos en punto de recogida.
• Falta de inventario visual frecuente en pasillos estrechos.

¿A quién afecta?
• Operadores, clientes y políticas de sostenibilidad local.

¿Por qué las soluciones actuales fallan?
• Sistemas monolíticos ROS+LLM difíciles de desplegar en el robot real.
• Poca separación entre “razonamiento del agente” y “ejecución ROS”.

¿Impacto si no se resuelve?
• Multas/incumplimiento ambiental, quiebres de stock, intervención humana constante.

¿Por qué ahora?
• LLM maduros + ROS 2 en edge; hace falta un marco ligero y repetible para dos tareas concretas en tienda.""",
        ],
    )

    # Slide 4
    apply_slide_texts(
        prs.slides[3],
        [
            "Figura sugerida: pasillo kombini + cubo de separación + estantería",
            "",
        ],
    )

    # Slide 5
    apply_slide_texts(
        prs.slides[4],
        [
            "",
            "Problema social: operación autónoma en espacio comercial reducido",
            """Use una figura: pasillo de kombini, cubo de separación (PET, latas, combustible) y estantería con productos.

Mensaje corto:
Los robots de servicio pueden ayudar en tareas repetitivas (clasificar residuo, registrar huecos en estantería), pero hoy integrar visión, lenguaje y movimiento en ROS suele requerir stacks pesados o acoplados al vendor. Se necesita un framework agnóstico al robot, con límites de seguridad y despliegue HTTP hacia el bridge ROS en el hardware real.""",
        ],
    )

    # Slide 6
    apply_slide_texts(
        prs.slides[5],
        ["Estado del arte: competidores y líneas base"],
    )

    # Slide 7
    apply_slide_texts(
        prs.slides[6],
        [
            "",
            "¿Por qué el estado del arte no es suficiente?",
            """Tres referencias directas (mismo espacio: lenguaje + ROS + robot real):

1. ROSA (NASA JPL) — agente conversacional ROS 1/2, amplio ecosistema académico.
2. ROS-LLM (Auromix) — integración rápida GPT + funciones ROS para navegación/movimiento.
3. ROSClaw — runtime physical-AI con daemon, permisos REAL y MCP; foco en seguridad corporal.

Criterios: código público, papers/demos, tareas language-conditioned, despliegue en robot.

OMX Framework no compite en “otro paper de navegación semántica”, sino en:
• Desacoplar agente (LangGraph ReAct + OpenAI) del robot vía REST.
• Perfiles JSON por robot (`configs/robots/`) sin reescribir el agente.
• Dos casos de uso kombini acotados para evidencia en vida real.""",
        ],
    )

    # Slide 8 — ROSA (6 shapes)
    apply_slide_texts(
        prs.slides[7],
        [
            "ROSA — NASA JPL (ROS Agent)",
            "",
            "",
            "Agente de lenguaje natural que descubre e invoca capacidades ROS; ROS 1/2; arXiv 2410.06472.",
            "¿Qué?\nInteracción en lenguaje natural con sistemas ROS existentes.",
            "¿Cómo?\nHerramientas ROS generadas + LLM en el mismo entorno ROS.",
        ],
    )

    # Slide 9 — ROSA results/limitations
    apply_slide_texts(
        prs.slides[8],
        [
            "Resultados",
            "Demostraciones en manipulación, diagnóstico y tareas ROS diversas; reduce barrera para operadores no expertos.",
            "Limitaciones",
            "Acoplamiento fuerte al stack ROS del host del agente; menos énfasis en bridge remoto y perfiles hardware mínimos; no centrado en retail/kombini ni en contrato REST uniforme.",
        ],
    )

    # Slide 10 — ROS-LLM (6 shapes)
    apply_slide_texts(
        prs.slides[9],
        [
            "ROS-LLM — Auromix (ROS 2 Humble)",
            "",
            "",
            "GPT-4/ChatGPT + funciones ROS registradas en ~10 min; navegación y manipulación embodied.",
            "¿Qué?\nControl y diálogo con cualquier robot ROS.",
            "¿Cómo?\nLLM + callbacks ROS en el mismo stack del robot.",
        ],
    )

    # Slide 11 — ROS-LLM results/limitations
    apply_slide_texts(
        prs.slides[10],
        [
            "Resultados",
            "Integración rápida; tutoriales para móvil y manipulación; comunidad activa en ROS 2.",
            "Limitaciones",
            "Modelo de extensión = código ROS en el mismo entorno; menos separación agente/edge; seguridad y límites de movimiento dependen de cada integración manual; no perfil declarativo único para brazo OMX + base móvil kombini.",
        ],
    )

    # Slide 12 — matrix intro
    apply_slide_texts(
        prs.slides[11],
        [
            "Matriz comparativa (OMX vs estado del arte)",
            "",
            "Filas: bridge HTTP, perfil JSON, LangGraph, safety declarativa, caso residuos, caso estanterías.\n"
            "Columnas: ROSA | ROS-LLM | ROSClaw | OMX Framework.\n"
            "(Completar tabla en PowerPoint con ✓ / ○ / —)",
        ],
    )

    # Slide 13 — ROSClaw (6 shapes)
    apply_slide_texts(
        prs.slides[12],
        [
            "ROSClaw — runtime physical-AI",
            "",
            "",
            "Daemon rosclawd, REAL/SHADOW, perfiles de cuerpo, MCP, receipts de ejecución; foco seguridad.",
            "¿Qué?\nControl físico con frontera de privilegios y simulación.",
            "¿Cómo?\nAgente vía MCP; sin ROS directo desde el chat.",
        ],
    )

    # Slide 14 — Gap
    apply_slide_texts(
        prs.slides[13],
        [
            "Brecha de investigación (research gap)",
            "",
            """Cómo identificar la brecha:
1. Competidores resuelven bien “ROS + LLM”, pero con distintos trade-offs (acoplamiento, seguridad pesada, o generalidad extrema).
2. La matriz muestra falta de un marco ligero: agente en laptop/cloud, robot en edge con bridge REST, dos pipelines de percepción para kombini.
3. La brecha importa porque el despliegue real en tienda exige iterar visión y movimiento sin recompilar todo el stack cada semana.

Brecha:
Falta un framework ROS-agente minimalista, reproducible y robot-agnóstico, optimizado para (A) clasificación de residuos/objetos en punto kombini y (B) navegación + análisis de estanterías, con límites de seguridad declarativos.""",
        ],
    )

    # Slide 15 — Gap example OMX (5 shapes)
    apply_slide_texts(
        prs.slides[14],
        [
            "Ejemplo brecha — OMX Framework",
            "",
            """Qué existe:
ROSA/ROS-LLM/ROSClaw cubren lenguaje + ROS con distintos niveles de seguridad y acoplamiento.

Qué falta:
Una línea directa Chat UI → LangGraph → tools desde JSON → bridge HTTP → ROS 2 en el robot OMX / base móvil, con prompts y tools específicos para residuos y shelf-audit.

Por qué importa:
Permite pilotos en kombini real con el mismo código del agente cambiando solo `configs/robots/*.json` y el bridge (`omx_bridge.py`, futuro `kombini_mobile_bridge`).""",
            "Research Gap:",
            "Un framework desacoplado que una clasificación de residuos y auditoría de estanterías con métricas de tienda (precisión, cobertura, tiempo por pasillo).",
        ],
    )

    # Slide 16
    apply_slide_texts(prs.slides[15], ["Objetivo de investigación y preguntas (RQs)"])

    # Slide 17 — Objective
    apply_slide_texts(
        prs.slides[16],
        [
            "Objetivo",
            "",
            """Objetivo principal:
Desarrollar y validar OMX Framework — agente LangGraph ReAct con OpenAI conectado a ROS 2 solo vía bridge HTTP y perfiles JSON — para dos aplicaciones en kombini real:
(1) clasificación y deposición asistida de residuos/objetos,
(2) navegación autónoma en pasillo con análisis de estanterías (huecos, desorden, etiquetas).

Debe diferenciarse por desacoplamiento operativo (agente sin ROS local), configuración declarativa y foco en evidencia de piloto retail, no en runtime de seguridad industrial completo.""",
        ],
    )

    # Slide 18 — RQs
    apply_slide_texts(
        prs.slides[17],
        [
            "Preguntas de investigación",
            "",
            """RQ1 — Clasificación de residuos:
¿El agente con visión + tools de brazo/gripper alcanza precisión y tiempo aceptables vs operador en cubo kombini?
→ Métricas: accuracy por clase (PET, lata, combustible, otros), tiempo por objeto, tasa de reintento.

RQ2 — Auditoría de estantería:
¿La navegación (cmd_vel / futuro base) + capturas programadas cubren ≥X% de frente de estantería por pasillo?
→ Métricas: cobertura (%), detección de hueco/OOS, tiempo por pasillo de 10 m.

RQ3 — Usabilidad del framework:
¿Cambiar de turtlesim → OMX-F → móvil kombini solo con JSON + bridge reduce horas de integración vs fork ROS-LLM?
→ Métrica: horas de integración, líneas de código nuevas.

RQ4 — Seguridad operativa:
¿Los límites en `safety.py` + perfil JSON evitan comandos fuera de rango en pruebas reales?
→ Métricas: violaciones bloqueadas, incidentes (0 deseado).""",
        ],
    )

    # Slide 19
    apply_slide_texts(prs.slides[18], ["Enfoque propuesto / metodología"])

    # Slide 20 — Proposal
    apply_slide_texts(
        prs.slides[19],
        [
            "Propuesta — OMX Framework",
            "",
            """Idea central:
Capas fijas: Chat UI (:8501) → `agent/graph.py` (ReAct) → `tools_lc.py` (desde perfil) → `adapter/http_bridge.py` → `bridge/*_bridge.py` → ROS 2.

Componentes nuevos para kombini:
• Módulo visión (detector residuo / producto) expuesto como tool o pre-proceso en bridge.
• Perfil `kombini_mobile.json` (navegación) + `open_manipulator_omx_f.json` (manipulación).
• Prompts en `prompts.py` con políticas de tienda (no bloquear pasillo, confirmar clase antes de soltar).

Novedad vs SOTA:
Desacoplamiento HTTP + perfiles; no sustituir rosclawd ni reimplementar ROSA; foco en dos KPIs de tienda.

Evaluación:
Baseline: operador manual + script ROS sin LLM; comparar con agente OMX en mismos escenarios.""",
        ],
    )

    # Slide 21 — diagram placeholder
    apply_slide_texts(
        prs.slides[20],
        [
            "Arquitectura OMX (añadir captura del README)",
            "Chat UI ──HTTP──► LangGraph ReAct ──HTTP──► Bridge REST ──► ROS 2\n"
            "                         ▲\n"
            "                 configs/robots/*.json",
        ],
    )

    # Slide 22
    apply_slide_texts(prs.slides[21], ["Sistema / configuración experimental"])

    # Slide 23 — Architecture
    apply_slide_texts(
        prs.slides[22],
        [
            "Arquitectura del sistema",
            "",
            """Hardware:
• OpenMANIPULATOR OMX-F (brazo) + cámara RGB(D) en muñeca o torso.
• Base móvil differential-drive (planeada) para pasillo kombini — o teleop inicial.
• PC edge en robot (Docker ROS 2 Jazzy) + laptop para agente.

Software:
• OMX: Python 3.10+, LangGraph, OpenAI (`gpt-4o-mini` default).
• Bridges: `omx_bridge.py`, `turtlesim_bridge.py` (sanity), futuro bridge móvil.
• Visión: modelo clasificación residuos (entrenar/fine-tune con datos tienda).

Escenarios:
A) Cubo de residuos junto a caja — pick, classify, deposit.
B) Un pasillo — waypoints, paradas, fotos de estantería, informe JSON.

Métricas ligadas a RQ1–RQ4.""",
            "Mensaje clave:\n"
            "Misma arquitectura agente/bridge para simulación (turtlesim) y piloto kombini real.",
        ],
    )

    # Slide 24
    apply_slide_texts(prs.slides[23], ["Trabajo completado desde la última actualización"])

    # Slide 25 — Progress
    apply_slide_texts(
        prs.slides[24],
        [
            f"Actualización: {TODAY}",
            "",
            """Logros concretos (repositorio omx-framework):
✓ Agente LangGraph ReAct con OpenAI y servidor chat HTTP.
✓ Adapter HTTP y contrato REST documentado (`bridge/README.md`).
✓ Bridge OMX-F para ROS 2 Jazzy (`omx_bridge.py`) + demo turtlesim Docker.
✓ Perfiles JSON (`open_manipulator_omx_f.json`, `generic_arm.json`, `turtlesim.json`).
✓ Capa `safety.py` con límites por paso de articulaciones.
✓ Tests unitarios de seguridad (`tests/test_safety.py`).

Pendiente hacia kombini:
○ Tool de visión residuos / shelf-audit
○ Bridge base móvil y waypoints
○ Dataset y experimentos en entorno tipo kombini""",
            "Mensaje clave:\n"
            "Base del framework lista; siguiente hito = piloto caso (1) clasificación de residuos.",
        ],
    )

    # Slide 26 — Issues
    apply_slide_texts(
        prs.slides[25],
        [
            "Problemas abiertos y riesgos",
            "",
            """1) Visión en tienda real
Problema: reflejos en envases PET y etiquetas pequeñas.
Causa: iluminación kombini + ángulo cámara.
Probado: ajuste umbral en simulación.
Siguiente: recopilar 200+ imágenes locales, fine-tune detector.

2) Latencia LLM + movimiento
Problema: ReAct multi-paso lento para objetos en cinta rápida.
Causa: iteraciones OpenAI + planificación secuencial.
Probado: reducir `ROS_AGENT_MAX_ITERATIONS`, cache de clases frecuentes.
Siguiente: modo “clasificar una vez” con tool única.

3) Base móvil no integrada aún
Problema: RQ2 no evaluable en pasillo real.
Causa: solo brazo en bridge actual.
Siguiente: definir REST `/cmd_vel` o `/navigate` en nuevo bridge.

4) Seguridad en espacio público
Problema: brazo cerca de clientes.
Siguiente: zonas prohibidas en JSON + velocidad máxima + E-stop hardware.""",
        ],
    )

    # Slide 27
    apply_slide_texts(
        prs.slides[26],
        [
            "Trabajo para las próximas semanas",
            "",
            """Objetivo próximo periodo (ej. 4–6 semanas):

Semana 1–2:
• Prototipo tool `classify_waste` + integración cámara en bridge.
• Guion de demo caso (1) en banco de pruebas (sin público).

Semana 3–4:
• Esqueleto `kombini_mobile_bridge` + waypoints en mapa reducido.
• Captura sistemática estantería → JSON de huecos.

Semana 5–6:
• 20+ trials residuos; tabla resultados RQ1.
• Primer pasillo con cobertura % (RQ2 preliminar).
• Actualizar README y slide de matriz con números reales.""",
        ],
    )

    # Slide 28 — Timeline
    apply_slide_texts(
        prs.slides[27],
        [
            "Cronograma y entregables esperados",
            "",
            """✓ Completado: framework agente + bridge OMX + turtlesim + safety
● Actual: diseño casos kombini + visión
○ Siguiente: piloto residuos → piloto estanterías → paper/informe técnico

Hitos:
• M1 — Demo clasificación residuo (bench) — entregable: video + métricas
• M2 — Navegación pasillo + shelf JSON — entregable: mapa cobertura
• M3 — Comparativa vs baseline manual — entregable: tablas RQ1–RQ3
• M4 — Documentación despliegue tienda — entregable: guía operador

Entregable final:
OMX Framework como producto open-source diferenciado (ligero, HTTP, JSON) con evidencia en dos flujos kombini reales.""",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    prs.save(str(OUT_ALT))
    print(f"Saved:\n  {OUT}\n  {OUT_ALT}")


if __name__ == "__main__":
    main()
